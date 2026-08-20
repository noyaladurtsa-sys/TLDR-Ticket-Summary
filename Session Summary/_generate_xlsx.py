from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import date

wb = Workbook()
ws = wb.active
ws.title = "Session Summary"

BRAND = "1B3A6B"
WHITE = "FFFFFF"
BORDER_COLOR = "D1D5DB"
ALT_ROW = "F8FAFC"

hdr_font = Font(name="Segoe UI", size=10, bold=True, color=WHITE)
hdr_fill = PatternFill("solid", fgColor=BRAND)
hdr_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

data_font = Font(name="Segoe UI", size=10)
data_align = Alignment(horizontal="left", vertical="top", wrap_text=True)
date_align = Alignment(horizontal="center", vertical="top", wrap_text=True)
alt_fill = PatternFill("solid", fgColor=ALT_ROW)
thin_border = Border(
    left=Side(style="thin", color=BORDER_COLOR),
    right=Side(style="thin", color=BORDER_COLOR),
    top=Side(style="thin", color=BORDER_COLOR),
    bottom=Side(style="thin", color=BORDER_COLOR),
)

headers = ["WE (Week Ending)", "What We Did", "Challenges", "Fix", "Test Ticket(s)", "Next Steps"]
for col, h in enumerate(headers, 1):
    c = ws.cell(row=1, column=col, value=h)
    c.font = hdr_font
    c.fill = hdr_fill
    c.alignment = hdr_align
    c.border = thin_border

D = date
rows = [
    [D(2026,5,25),
     "Created Streamlit app with HappyFox API integration\nBuilt Claude AI analysis pipeline (structured JSON output)\nSet up PostgreSQL DB (metaboard) with caching\nBuilt HTML report generator with download/open/print\nCreated View Saved Reports page with token cost tracking\nPushed initial commit to GitHub",
     "HappyFox API returns updates in two different formats (embedded vs endpoint) with different field names",
     "Built dual-format parser that checks both structures and normalizes to a common format",
     "#29640 - MX4200 (early test, no updates)\n#130211 - MX6200 (early test, no updates)\n#129157 - FGW5500 (early test, no updates)\n#129347 - MX6200 (early test, no updates)\n#128936 - FGMM1000 (early test, no updates)\n#130546 - MX4200 Overload (19-25 updates, tested 4x)\n#129512 - MBE7000 (72 updates)\n#120920 - LN1400/MBE7000 (122-126 updates)\n#13054 - Chat conversation (9 updates)",
     "Design QA scoring system\nExpand Claude prompt for deeper analysis"],

    [D(2026,5,25),
     "",
     "API credentials must not be committed to git",
     "Used Streamlit's built-in secrets.toml system and added it to .gitignore",
     "",
     ""],

    [D(2026,5,29),
     "Continued testing with more ticket types\nTested non-English tickets (Spanish)\nTested high-update-count tickets (100+ updates)\nRefined prompt and output formatting",
     "Large tickets with 100+ updates hit token limits and took longer to process",
     "Monitored token usage per ticket; prioritised caching to avoid repeat API costs",
     "#130381 - FGW5500 (28 updates)\n#129640 - E8450 WiFi (56 updates)\n#131007 - FGW5500 (13 updates)\n#121534 - FGW5500 (246 updates - largest)\n#26188 - Un nodo sin conexion (7 updates, Spanish)\n#126188 - MBE7002 Wired Disconnection (193 updates)\n#109928 - FGW5500 (96 updates)",
     "Build QA scoring system\nAdd agent-level analysis"],

    [D(2026,6,7),
     "Tested with IPOE connection ticket\nValidated prompt against complex multi-agent ticket (152 updates)",
     "Ticket #127911 with 152 updates and 21K input tokens - highest token usage so far",
     "Caching ensures repeat views cost nothing; tracked token usage in saved reports page",
     "#127911 - LN6001 IPOE connection error (152 updates, 21K input tokens)",
     "Redesign Claude prompt for QA scoring"],

    [D(2026,6,14),
     "Redesigned Claude prompt (Senior QA Analyst + Linksys Technician + CX Director)\nBuilt dual scoring system (1-5): Technical Accuracy & Customer Handling per agent\nAdded Agent Scorecard table with roles, scores, flags, summaries\nAdded Coaching Notes (agent, finding, correct script, policy ref)\nAdded Coaching Priority Banner (URGENT / REVIEW / GOOD)\nAdded Technical QA Findings, Soft Skill Findings, Opportunities sections\nBuilt delete_cache.py utility\nExpanded JSON output to 30+ fields\nCreated feature branch and opened PR #1",
     "Claude returned scores as strings instead of integers, breaking scorecard rendering",
     "Added isinstance type checking for score comparisons and color-coding logic",
     "#131976 - MX6201 (56 updates)\n#132182 - MX6201 (47 updates)",
     "Implement Gemini audit recommendations\nBuild premature close detection"],

    [D(2026,6,14),
     "",
     "Agent names appeared as full email addresses (e.g., jorgenathaniel.amores@concentrix.com)",
     "Built _clean_agent_name() helper to strip @concentrix.com and @linksys.com domains",
     "",
     ""],

    [D(2026,6,14),
     "",
     "Anthropic API credits drained to $19.64 unexpectedly",
     "Discovered Claude Code CLI defaulting to Opus model. Locked app to claude-sonnet-4-6",
     "",
     ""],

    [D(2026,6,16),
     "Analysed Gemini Internal Auditing Memorandum (ticket #127492)\nBuilt Premature Close Detection (Observe and Close KPI trap)\nBuilt Callback Breach Tracking (promise vs actual follow-up)\nBuilt Complex Environment Checklist (5 checks for 4+ node mesh)\nBuilt session-grouped Detailed Case Summary with colored dots and inline flags\nReorganized UI into two-tab layout (Summary + Deep Dive)\nAdded auto-translation via deep-translator\nBuilt 4 Plotly visualizations (Dot Timeline, Sankey, Donut, Swimlane)\nBuilt HTML report equivalents for all new sections\nFixed report opening via base64 Blob URL",
     "Plotly Sankey rejected 8-digit hex colors (#d9770680)",
     "Built _hex_rgba() helper to convert to rgba(r,g,b,alpha) format",
     "#131478 - MR8300 slow speed (55 updates)\n#118250 - MX5500 separate 2.4GHz (20 updates)\n#132712 - FGW5500 (59 updates)\n#127492 - MX4200 No internet (170 updates) - primary QA test ticket",
     "Add Actions Taken section\nFix Agent Scorecard cache bug"],

    [D(2026,6,16),
     "",
     "Opening HTML report broke when it contained Plotly script tags",
     "Switched to base64-encoding HTML, creating Blob URL, opening in new tab",
     "",
     ""],

    [D(2026,6,16),
     "",
     "Dot Timeline clustered all entries on same date - unreadable",
     "Used entry index (#1, #2...) on X axis instead of actual dates",
     "",
     ""],

    [D(2026,6,16),
     "",
     "What Happened section was flat dump of 50+ raw updates with email addresses",
     "Built build_session_groups() to group by date gaps and key events, with session titles and cleaned names",
     "",
     ""],

    [D(2026,6,16),
     "",
     "Non-English messages (Spanish, Filipino) not readable by supervisors",
     "Added _auto_translate() using deep-translator GoogleTranslator (triggers when >30% non-ASCII)",
     "",
     ""],

    [D(2026,6,18),
     "Added Actions Taken section (collapsed expander, [Agent / Date] format)\nDiscovered Agent Scorecard cache bug - #127492 has empty agents_involved\nInvestigated DB and confirmed empty ai_json\nIdentified short test tickets (#133821, #133824, #133827 - 2 msgs each, ~$0.02)\nCreated session handoff system with passphrase resume",
     "Agent Scorecard stopped appearing - cached ai_json has empty agents_involved array",
     "Need to delete cached row with delete_cache.py and re-analyse ticket",
     "#127492 - MX4200 (cache investigation, scorecard bug)\nIdentified for next test:\n#133821 - Dropping connection (2 msgs, ~$0.02)\n#133824 - Router default name (2 msgs, ~$0.02)\n#133827 - Product features inquiry (2 msgs, ~$0.02)",
     "Test Actions Taken with short ticket (#133821)\nClear cache and re-run #127492 for scorecard\nSync HTML report layout to match two-tab UI\nCommit June 18 changes to git"],

    [D(2026,6,18),
     "",
     "Re-running #127492 is expensive (large ticket, many updates)",
     "Queried HappyFox API for short tickets. Found 2-message tickets costing ~$0.02 vs $0.10+",
     "",
     ""],
]

for i, r in enumerate(rows):
    row_num = i + 2
    ws.cell(row=row_num, column=1, value=r[0]).number_format = "DD/MM/YYYY"
    ws.cell(row=row_num, column=2, value=r[1])
    ws.cell(row=row_num, column=3, value=r[2])
    ws.cell(row=row_num, column=4, value=r[3])
    ws.cell(row=row_num, column=5, value=r[4])
    ws.cell(row=row_num, column=6, value=r[5])

    is_alt = (i % 2 == 1)
    for col in range(1, 7):
        cell = ws.cell(row=row_num, column=col)
        cell.font = data_font
        cell.alignment = date_align if col == 1 else data_align
        cell.border = thin_border
        if is_alt:
            cell.fill = alt_fill

col_widths = [15, 60, 50, 50, 45, 40]
for i, w in enumerate(col_widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w

ws.freeze_panes = "A2"

out = r"C:\Users\Ashley-Admin\Desktop\Claude\Project\Happy Link Claude TLDR Ticket Summary\Session Summary\TLDR_Session_Summary.xlsx"
wb.save(out)
print("Saved to " + out)
