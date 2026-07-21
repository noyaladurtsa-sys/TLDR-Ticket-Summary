import docx
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml
import os

BRAND = RGBColor(0x1B, 0x3A, 0x6B)
DARK_GRAY = RGBColor(0x37, 0x41, 0x51)
MID_GRAY = RGBColor(0x6B, 0x72, 0x80)
RED_ACCENT = RGBColor(0x99, 0x1B, 0x1B)
GREEN_ACCENT = RGBColor(0x16, 0x6A, 0x34)
AMBER_ACCENT = RGBColor(0x92, 0x40, 0x0E)

doc = Document()

style = doc.styles['Normal']
font = style.font
font.name = 'Segoe UI'
font.size = Pt(10)
font.color.rgb = DARK_GRAY

for section in doc.sections:
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

# Header
header = doc.sections[0].header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
run = hp.add_run("TLDR Ticket Summary — Session Summary")
run.font.size = Pt(9)
run.font.color.rgb = MID_GRAY
run.font.name = 'Segoe UI'
# Header bottom border
pPr = hp._p.get_or_add_pPr()
pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="4" w:space="4" w:color="1B3A6B"/></w:pBdr>')
pPr.append(pBdr)

# Footer with page numbers
footer = doc.sections[0].footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = fp.add_run("Page ")
run.font.size = Pt(8)
run.font.color.rgb = MID_GRAY
run.font.name = 'Segoe UI'
fld_xml = parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr=" PAGE "><w:r><w:t>1</w:t></w:r></w:fldSimple>')
fp._p.append(fld_xml)


def add_title(text, size=22, color=BRAND, space_after=4):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.space_after = Pt(space_after)
    p.space_before = Pt(0)
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.name = 'Segoe UI'
    return p


def add_subtitle(text, size=12, color=MID_GRAY):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.space_after = Pt(16)
    p.space_before = Pt(0)
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.name = 'Segoe UI'
    return p


def add_heading_styled(text, level=1):
    p = doc.add_paragraph()
    p.space_before = Pt(18 if level == 1 else 12)
    p.space_after = Pt(6)
    run = p.add_run(text)
    run.bold = True
    run.font.name = 'Segoe UI'
    if level == 1:
        run.font.size = Pt(16)
        run.font.color.rgb = BRAND
        pPr = p._p.get_or_add_pPr()
        pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="4" w:color="1B3A6B"/></w:pBdr>')
        pPr.append(pBdr)
    elif level == 2:
        run.font.size = Pt(13)
        run.font.color.rgb = BRAND
    elif level == 3:
        run.font.size = Pt(11)
        run.font.color.rgb = DARK_GRAY
    return p


def add_body(text, space_after=6):
    p = doc.add_paragraph()
    p.space_after = Pt(space_after)
    p.space_before = Pt(0)
    run = p.add_run(text)
    run.font.size = Pt(10)
    run.font.color.rgb = DARK_GRAY
    run.font.name = 'Segoe UI'
    return p


def add_bullet(text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.space_after = Pt(3)
    p.space_before = Pt(1)
    if bold_prefix:
        run = p.add_run(bold_prefix)
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = DARK_GRAY
        run.font.name = 'Segoe UI'
        run = p.add_run(text)
        run.font.size = Pt(10)
        run.font.color.rgb = DARK_GRAY
        run.font.name = 'Segoe UI'
    else:
        for run in p.runs:
            run.font.size = Pt(10)
            run.font.name = 'Segoe UI'
        if not p.runs:
            run = p.add_run(text)
            run.font.size = Pt(10)
            run.font.color.rgb = DARK_GRAY
            run.font.name = 'Segoe UI'
    return p


def add_challenge(challenge, fix):
    p = doc.add_paragraph(style='List Bullet')
    p.space_after = Pt(6)
    p.space_before = Pt(2)
    run = p.add_run("Challenge: ")
    run.bold = True
    run.font.size = Pt(10)
    run.font.color.rgb = RED_ACCENT
    run.font.name = 'Segoe UI'
    run = p.add_run(challenge + "\n")
    run.font.size = Pt(10)
    run.font.color.rgb = DARK_GRAY
    run.font.name = 'Segoe UI'
    run = p.add_run("Fix: ")
    run.bold = True
    run.font.size = Pt(10)
    run.font.color.rgb = GREEN_ACCENT
    run.font.name = 'Segoe UI'
    run = p.add_run(fix)
    run.font.size = Pt(10)
    run.font.color.rgb = DARK_GRAY
    run.font.name = 'Segoe UI'
    return p


def add_divider():
    p = doc.add_paragraph()
    p.space_before = Pt(4)
    p.space_after = Pt(4)
    pPr = p._p.get_or_add_pPr()
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="4" w:space="4" w:color="E5E7EB"/></w:pBdr>')
    pPr.append(pBdr)


# ============================================================
# TITLE PAGE
# ============================================================
doc.add_paragraph()  # spacer
add_title("TLDR Ticket Summary Project", size=26)
add_title("Weekly Session Summary", size=18, color=MID_GRAY, space_after=8)
add_subtitle("May 23 – June 18, 2026", size=13, color=MID_GRAY)
add_divider()

# PROJECT OVERVIEW
add_heading_styled("Project Overview", level=1)
add_body(
    "A Streamlit dashboard that fetches HappyFox support tickets, analyses them with Claude AI "
    "(QA scorecard + coaching insights), and saves executive reports to a local PostgreSQL database. "
    "Built for internal QA coaching at Linksys."
)
add_body(
    "The app allows supervisors to enter any ticket number, receive an AI-generated executive summary "
    "with dual agent scoring, coaching notes, and systemic failure detection — then save, download, "
    "or print the report. What previously took 15–30 minutes of manual review now takes seconds."
)

# ============================================================
# WEEK 1
# ============================================================
doc.add_page_break()
add_heading_styled("Week 1: May 23 – May 25, 2026", level=1)
add_heading_styled("Session: Project Initialization", level=2)

add_heading_styled("What We Did", level=3)
add_bullet("Created the initial Streamlit app (app.py) with HappyFox API integration to fetch live support tickets")
add_bullet("Built Claude AI analysis pipeline — sends ticket text to Claude Sonnet, receives structured JSON with problem summary, categories, root cause, and recommended next steps")
add_bullet("Set up PostgreSQL database (metaboard) with executive_reports table for caching analysed tickets (zero API cost on repeat views)")
add_bullet("Built HTML executive report generator with download, open-in-new-tab, and print functionality")
add_bullet("Created the View Saved Reports page (pages/1_View_Saved_Reports.py) with token usage and cost tracking")
add_bullet("Pushed initial commit to GitHub (noyaladurtsa-sys/TLDR-Ticket-Summary)")

add_heading_styled("What It’s For", level=3)
add_body(
    "The foundation of the tool — allows Ashley to enter any HappyFox ticket number, get an "
    "AI-generated executive summary with QA insights, and save/download reports. Replaces manual "
    "ticket review which took 15–30 minutes per ticket."
)

add_heading_styled("Challenges & Fixes", level=3)
add_challenge(
    "HappyFox API returns updates in two different formats (embedded vs endpoint) with different "
    "field names (by.name vs agent.name, timestamp vs created_at, message.text vs text).",
    "Built dual-format parser that checks both structures and normalizes to a common format."
)
add_challenge(
    "API credentials management — secrets must not be committed to git.",
    "Used Streamlit’s built-in secrets.toml system and added it to .gitignore."
)

# ============================================================
# WEEK 2
# ============================================================
doc.add_page_break()
add_heading_styled("Week 2: June 9 – June 14, 2026", level=1)
add_heading_styled("Session: QA Scorecard & Coaching System", level=2)

add_heading_styled("What We Did", level=3)
add_bullet("Redesigned the Claude prompt to act as Senior QA Analyst + Expert Linksys Technician + Customer Experience Director")
add_bullet("Built dual scoring system (1–5 scale): Technical Accuracy and Customer Handling scores per agent")
add_bullet("Added Agent Scorecard table — lists every agent involved with their role, tech score, handling score, flags (FCR Violation, Callback Breach, Premature Close), and contribution summary")
add_bullet("Added Coaching Notes section — for each finding: agent name, what they did wrong, correct script, and policy reference")
add_bullet("Added Coaching Priority Banner — color-coded URGENT (red), REVIEW (yellow), GOOD (green)")
add_bullet("Added Technical QA Findings and Soft Skill Findings as separate sections")
add_bullet("Added Opportunities section with combined QA findings")
add_bullet("Built delete_cache.py utility to clear cached reports for re-testing after prompt changes")
add_bullet("Expanded JSON output from Claude to 30+ fields including model_number, serial_number, warranty_status, customer_phone, customer_country")
add_bullet("Created feature/qa-scorecard-sections branch and opened PR #1 against master")

add_heading_styled("What It’s For", level=3)
add_body(
    "Transforms the app from a basic ticket summarizer into a full QA coaching tool. Supervisors can "
    "now see exactly which agents need training, on what topics, with specific scripts they should use. "
    "The dual scoring system separates technical competence from customer handling — an agent might "
    "be technically strong but poor at empathy, or vice versa."
)

add_heading_styled("Challenges & Fixes", level=3)
add_challenge(
    "Claude sometimes returned scores as strings (\"3\") instead of integers, causing rendering issues in the scorecard.",
    "Added type checking (isinstance) for score comparisons and color-coding logic."
)
add_challenge(
    "Agent names appeared as full email addresses (e.g., jorgenathaniel.amores@concentrix.com).",
    "Built _clean_agent_name() helper to strip @concentrix.com and @linksys.com domains to readable names."
)
add_challenge(
    "Anthropic API credits drained to $19.64 unexpectedly.",
    "Discovered Claude Code CLI was defaulting to Opus model instead of Sonnet. Locked the app to claude-sonnet-4-6 and added cost awareness to workflow."
)

# ============================================================
# WEEK 3
# ============================================================
doc.add_page_break()
add_heading_styled("Week 3: June 15 – June 16, 2026", level=1)
add_heading_styled("Session: QA Deep Dive Features (Gemini Audit Implementation)", level=2)

add_heading_styled("What We Did", level=3)
add_bullet("Analysed Gemini’s Internal Auditing Memorandum PDF (based on ticket #127492, Richard Stell, MX4200) which identified three pillars of improvement")
add_bullet("Built Premature Close Detection — scans ticket lifecycle for agents closing tickets without confirmed fixes (\"Observe and Close\" KPI trap). Red/orange cards showing agent, date, what happened, and pattern type")
add_bullet("Built Callback Breach Tracking — scans every agent message for callback promises, cross-references whether follow-up happened within promised window. Table with green (Honored) and red (Breached) rows")
add_bullet("Built Complex Environment Checklist — for tickets with 4+ mesh nodes or mixed hardware, evaluates 5 diagnostic checks: Node Inventory, RSSI/Signal Check, Wired Bypass Test, Topology Reconciliation, Firmware Consistency")
add_bullet("Built session-grouped Detailed Case Summary — groups 50+ raw updates into logical sessions by date gaps and key events, with colored dots, inline flag badges, and summary bar")
add_bullet("Reorganized the entire UI into two-tab layout: Summary tab and Deep Dive tab")
add_bullet("Added auto-translation for non-English messages using deep-translator (GoogleTranslator)")
add_bullet("Built 4 interactive Plotly visualizations: Dot Timeline, Sankey diagram (handoffs), Donut chart (message distribution), Swimlane chart (critical events)")
add_bullet("Built HTML report equivalents for all new sections")
add_bullet("Fixed report opening mechanism — switched from document.write() to base64 Blob URL approach")

add_heading_styled("What It’s For", level=3)
add_body(
    "Implements the three pillars from the Gemini audit directly into the app. Now the tool doesn’t "
    "just score agents — it catches specific systemic failures: premature closures that inflate "
    "resolution metrics, broken callback promises that erode customer trust, and incomplete diagnostics "
    "on complex environments. The session-grouped view turns an unreadable wall of 50+ updates into a "
    "structured narrative with flagged critical events."
)

add_heading_styled("Challenges & Fixes", level=3)
add_challenge(
    "Plotly Sankey diagram rejected 8-digit hex colors (#d9770680).",
    "Built _hex_rgba() helper to convert to rgba(r,g,b,alpha) format which Sankey accepts."
)
add_challenge(
    "Opening HTML report in new tab broke when report contained Plotly script tags — document.write() couldn’t handle them.",
    "Switched to base64-encoding the HTML, creating a Blob URL, and opening that in a new tab."
)
add_challenge(
    "Dot Timeline chart clustered all entries on the same date making it unreadable.",
    "Used entry index (#1, #2…) on X axis instead of actual dates, with dates shown as labels."
)
add_challenge(
    "\"What Happened?\" section was a flat dump of 50+ raw updates with email addresses — unreadable.",
    "Built build_session_groups() to group by date gaps (2+ days) and key events (reopens, escalations), with auto-generated session titles and cleaned agent names."
)
add_challenge(
    "Non-English ticket messages (Spanish, Filipino) were not readable by English-speaking supervisors.",
    "Added _auto_translate() using deep-translator GoogleTranslator, triggered when >30% of characters are non-ASCII."
)

# ============================================================
# WEEK 4
# ============================================================
doc.add_page_break()
add_heading_styled("Week 4: June 17 – June 18, 2026", level=1)
add_heading_styled("Session: Actions Taken & Debugging", level=2)

add_heading_styled("What We Did", level=3)
add_bullet("Added Actions Taken section — collapsed expander between Summary and Coaching Priority Banner showing concise action log entries in [Agent · Date] format")
add_bullet("Discovered Agent Scorecard cache bug — ticket #127492 cached on June 16 has empty agents_involved array, causing the scorecard to not render")
add_bullet("Investigated database and confirmed ai_json for #127492 contains agents_involved: [] (empty)")
add_bullet("Identified short test tickets to save API credits: #133821 (2 messages), #133824 (2 messages), #133827 (2 messages) — estimated cost ~$0.02 each vs $0.10+ for large tickets")
add_bullet("Created session handoff system with passphrase-based resume pattern")

add_heading_styled("What It’s For", level=3)
add_body(
    "The Actions Taken section gives reviewers a quick chronological view of everything that happened "
    "in the ticket without expanding the full Deep Dive — a \"what did we do\" at a glance. The cache "
    "investigation explains why the Agent Scorecard disappeared and identified the fix (clear cache, "
    "re-run analysis)."
)

add_heading_styled("Challenges & Fixes", level=3)
add_challenge(
    "Agent Scorecard stopped appearing in the UI for ticket #127492.",
    "Investigated the database — found the cached ai_json from June 16 has an empty agents_involved "
    "array. The scorecard code correctly skips rendering when empty. Solution: delete the cached row "
    "with delete_cache.py and re-analyse the ticket to regenerate with populated agents_involved data."
)
add_challenge(
    "Re-running ticket #127492 is expensive (large ticket, many updates).",
    "Queried HappyFox API to find recent tickets sorted by message count. Identified 2-message tickets "
    "(#133821, #133824, #133827) that cost ~$0.02 per analysis instead of $0.10+."
)

# ============================================================
# OPEN ITEMS TABLE
# ============================================================
doc.add_page_break()
add_heading_styled("Open Items", level=1)

open_items = [
    ("Agent Scorecard for #127492", "Pending", "Need to clear cache and re-analyse"),
    ("Actions Taken section", "Untested", "Code added, needs live verification"),
    ("HTML report layout sync", "Not Started", "HTML report still uses old single-section layout, needs updating to match two-tab UI"),
    ("Short ticket test run", "Pending", "Run #133821 to verify all sections render correctly"),
    ("Git commit for June 18 changes", "Pending", "Actions Taken code not yet committed"),
]

table = doc.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = True

# Header row
hdr = table.rows[0]
for i, text in enumerate(["Item", "Status", "Notes"]):
    cell = hdr.cells[i]
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    run.font.name = 'Segoe UI'
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="1B3A6B" w:val="clear"/>')
    cell._tc.get_or_add_tcPr().append(shading)

# Data rows
status_colors = {
    "Pending": RGBColor(0x92, 0x40, 0x0E),
    "Untested": RGBColor(0x92, 0x40, 0x0E),
    "Not Started": RGBColor(0x99, 0x1B, 0x1B),
}

for item, status, notes in open_items:
    row = table.add_row()
    # Item
    c0 = row.cells[0]
    c0.text = ""
    run = c0.paragraphs[0].add_run(item)
    run.font.size = Pt(10)
    run.font.name = 'Segoe UI'
    run.font.color.rgb = DARK_GRAY
    run.bold = True

    # Status
    c1 = row.cells[1]
    c1.text = ""
    run = c1.paragraphs[0].add_run(status)
    run.font.size = Pt(10)
    run.font.name = 'Segoe UI'
    run.font.color.rgb = status_colors.get(status, DARK_GRAY)
    run.bold = True

    # Notes
    c2 = row.cells[2]
    c2.text = ""
    run = c2.paragraphs[0].add_run(notes)
    run.font.size = Pt(10)
    run.font.name = 'Segoe UI'
    run.font.color.rgb = DARK_GRAY

# Style table borders
for row in table.rows:
    for cell in row.cells:
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        borders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            '<w:top w:val="single" w:sz="4" w:space="0" w:color="D1D5DB"/>'
            '<w:left w:val="single" w:sz="4" w:space="0" w:color="D1D5DB"/>'
            '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="D1D5DB"/>'
            '<w:right w:val="single" w:sz="4" w:space="0" w:color="D1D5DB"/>'
            '</w:tcBorders>'
        )
        tcPr.append(borders)
        # Add cell padding
        mar = parse_xml(
            f'<w:tcMar {nsdecls("w")}>'
            '<w:top w:w="60" w:type="dxa"/>'
            '<w:left w:w="100" w:type="dxa"/>'
            '<w:bottom w:w="60" w:type="dxa"/>'
            '<w:right w:w="100" w:type="dxa"/>'
            '</w:tcMar>'
        )
        tcPr.append(mar)

# Set column widths
for row in table.rows:
    row.cells[0].width = Cm(5)
    row.cells[1].width = Cm(2.5)
    row.cells[2].width = Cm(9)

# ============================================================
# SAVE
# ============================================================
out_path = r"C:\Users\Ashley-Admin\Desktop\Claude\Project\Happy Link Claude TLDR Ticket Summary\Session Summary\TLDR_Session_Summary.docx"
doc.save(out_path)
print(f"Saved to {out_path}")
