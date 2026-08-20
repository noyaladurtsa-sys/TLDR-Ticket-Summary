from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = Workbook()
ws = wb.active
ws.title = "Transfer Guide"

BRAND = "1B3A6B"
WHITE = "FFFFFF"
BORDER_COLOR = "D1D5DB"
SECTION_BG = "F1F5F9"
GREEN = "16a34a"

hdr_font = Font(name="Segoe UI", size=11, bold=True, color=WHITE)
hdr_fill = PatternFill("solid", fgColor=BRAND)
hdr_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

title_font = Font(name="Segoe UI", size=14, bold=True, color=BRAND)
subtitle_font = Font(name="Segoe UI", size=10, color="6b7280")
section_font = Font(name="Segoe UI", size=10, bold=True, color=BRAND)
section_fill = PatternFill("solid", fgColor=SECTION_BG)

data_font = Font(name="Segoe UI", size=10)
data_align = Alignment(horizontal="left", vertical="top", wrap_text=True)
center_align = Alignment(horizontal="center", vertical="top", wrap_text=True)
done_font = Font(name="Segoe UI", size=10, bold=True, color=GREEN)

thin_border = Border(
    left=Side(style="thin", color=BORDER_COLOR),
    right=Side(style="thin", color=BORDER_COLOR),
    top=Side(style="thin", color=BORDER_COLOR),
    bottom=Side(style="thin", color=BORDER_COLOR),
)

col_widths = [8, 30, 75, 10]
for i, w in enumerate(col_widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w

# Title row
ws.merge_cells("A1:D1")
c = ws.cell(row=1, column=1, value="TLDR Ticket Summary - Project Transfer Guide")
c.font = title_font
c.alignment = Alignment(horizontal="left", vertical="center")

# Subtitle row
ws.merge_cells("A2:D2")
c = ws.cell(row=2, column=1, value="Transfer from current PC to new computer  |  Generated: 16 July 2026")
c.font = subtitle_font
c.alignment = Alignment(horizontal="left", vertical="center")

ws.row_dimensions[1].height = 30
ws.row_dimensions[2].height = 20

# Header row
headers = ["Step", "Task", "Command / Details", "Status"]
for col, h in enumerate(headers, 1):
    c = ws.cell(row=3, column=col, value=h)
    c.font = hdr_font
    c.fill = hdr_fill
    c.alignment = hdr_align
    c.border = thin_border

def add_section(row, text):
    ws.merge_cells(f"A{row}:D{row}")
    c = ws.cell(row=row, column=1, value=text)
    c.font = section_font
    c.fill = section_fill
    c.alignment = Alignment(horizontal="left", vertical="center")
    c.border = thin_border
    for col in range(2, 5):
        cell = ws.cell(row=row, column=col)
        cell.fill = section_fill
        cell.border = thin_border
    ws.row_dimensions[row].height = 22

def add_step(row, step, task, details, status=""):
    ws.cell(row=row, column=1, value=step).alignment = center_align
    ws.cell(row=row, column=2, value=task).alignment = data_align
    ws.cell(row=row, column=3, value=details).alignment = data_align
    sc = ws.cell(row=row, column=4, value=status)
    sc.alignment = center_align
    if status == "Done":
        sc.font = done_font
    else:
        sc.font = data_font
    for col in range(1, 5):
        cell = ws.cell(row=row, column=col)
        if cell.font == Font():
            cell.font = data_font
        cell.border = thin_border

r = 4

add_section(r, "BEFORE YOU START (ON CURRENT PC)")
r += 1

add_step(r, 1, "Push code to GitHub",
         "Already done (commit f8cb1da pushed to feature/qa-scorecard-sections on 16 Jul 2026)", "Done")
r += 1

add_step(r, 2, "Export database backup",
         "Already done (28 reports exported to executive_reports_backup.csv, pushed to GitHub)", "Done")
r += 1

add_step(r, 3, "Note your secrets",
         "Copy .streamlit/secrets.toml to a secure location (USB or cloud storage).\n"
         "This file contains API keys and is NOT in git.\n"
         "See the Secrets Template section below for the file structure.", "")
r += 1

add_section(r, "INSTALL SOFTWARE (ON NEW PC)")
r += 1

add_step(r, 4, "Install Python 3.13+",
         'Download from https://python.org\nIMPORTANT: Check "Add Python to PATH" during installation', "")
r += 1

add_step(r, 5, "Install PostgreSQL",
         "Download from https://postgresql.org\nDuring install, set postgres user password to: 5432\n"
         "Default port: 5432", "")
r += 1

add_step(r, 6, "Install Git",
         "Download from https://git-scm.com\nDefault settings are fine", "")
r += 1

add_section(r, "SET UP PROJECT (ON NEW PC)")
r += 1

add_step(r, 7, "Clone the repository",
         "git clone https://github.com/noyaladurtsa-sys/TLDR-Ticket-Summary.git", "")
r += 1

add_step(r, 8, "Switch to feature branch",
         "cd TLDR-Ticket-Summary\ngit checkout feature/qa-scorecard-sections", "")
r += 1

add_step(r, 9, "Install Python packages",
         "pip install streamlit anthropic psycopg2-binary plotly deep-translator\n\n"
         "Exact versions (optional):\n"
         "pip install streamlit==1.57.0 anthropic==0.102.0 psycopg2-binary==2.9.12 plotly==6.7.0 deep-translator==1.9.1", "")
r += 1

add_step(r, 10, "Create the database",
         "Open a terminal and run:\ncreatedb -U postgres metaboard\n\n"
         "If createdb is not found, add PostgreSQL bin folder to your PATH:\n"
         r'e.g. C:\Program Files\PostgreSQL\13\bin', "")
r += 1

add_step(r, 11, "Restore cached reports",
         "python _restore_db.py\n\n"
         "This restores 28 cached reports from CSV backup.\n"
         "Cached tickets load instantly with zero API cost.", "")
r += 1

add_step(r, 12, "Create secrets file",
         "Create a folder called .streamlit inside the project folder.\n"
         "Create a file called secrets.toml inside .streamlit/\n"
         "Paste the contents from Step 3.\n"
         "See the Secrets Template section below.", "")
r += 1

add_section(r, "VERIFY")
r += 1

add_step(r, 13, "Run the app",
         "streamlit run app.py\n\nThe app should open in your browser at http://localhost:8501", "")
r += 1

add_step(r, 14, "Test cached ticket",
         "Enter ticket #127492\nShould load from DB instantly with zero API cost.\n"
         "Check: Ticket Overview, Customer Details, The Problem, Summary, Actions Taken all render.", "")
r += 1

add_step(r, 15, "Test fresh ticket",
         "Enter ticket #133821 (2 messages, ~$0.02 API cost)\n"
         "Verifies: API connection works, all sections render, Agent Scorecard appears with scores.", "")
r += 1

# Secrets template section
r += 1
add_section(r, "SECRETS FILE TEMPLATE")
r += 1

ws.merge_cells(f"A{r}:D{r}")
note = ws.cell(row=r, column=1,
    value='Create .streamlit/secrets.toml with this structure (fill in your actual values from the current PC):')
note.font = Font(name="Segoe UI", size=10, italic=True, color="6b7280")
note.alignment = data_align
note.border = thin_border
for col in range(2, 5):
    ws.cell(row=r, column=col).border = thin_border
r += 1

secret_headers = ["", "Key", "Value", ""]
for col, h in enumerate(secret_headers, 1):
    c = ws.cell(row=r, column=col, value=h)
    c.font = Font(name="Segoe UI", size=10, bold=True, color=WHITE) if h else data_font
    c.fill = hdr_fill if h else PatternFill()
    c.alignment = center_align
    c.border = thin_border
r += 1

secrets = [
    ("HAPPYFOX_BASE_URL", "https://support.linksys.com"),
    ("HAPPYFOX_API_KEY", "(copy from current secrets.toml)"),
    ("HAPPYFOX_AUTH_CODE", "(copy from current secrets.toml)"),
    ("ANTHROPIC_API_KEY", "(copy from current secrets.toml)"),
    ("DB_HOST", "localhost"),
    ("DB_PORT", "5432"),
    ("DB_NAME", "metaboard"),
    ("DB_USER", "postgres"),
    ("DB_PASS", "5432"),
]

for key, val in secrets:
    ws.cell(row=r, column=1, value="").border = thin_border
    k = ws.cell(row=r, column=2, value=key)
    k.font = Font(name="Segoe UI", size=10, bold=True)
    k.alignment = data_align
    k.border = thin_border
    v = ws.cell(row=r, column=3, value=val)
    v.font = data_font
    v.alignment = data_align
    v.border = thin_border
    if "(copy" in val:
        v.font = Font(name="Segoe UI", size=10, italic=True, color="9a3412")
    ws.cell(row=r, column=4, value="").border = thin_border
    r += 1

ws.freeze_panes = "A4"

out = r"C:\Users\Ashley-Admin\Desktop\Claude\Project\Happy Link Claude TLDR Ticket Summary\Session Summary\TLDR_Transfer_Guide.xlsx"
wb.save(out)
print("Saved to " + out)
